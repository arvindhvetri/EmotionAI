from pptx import Presentation
import sys

def extract_text_from_ppt(ppt_path):
    prs = Presentation(ppt_path)
    for i, slide in enumerate(prs.slides):
        print(f"--- Slide {i+1} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                print(shape.text.encode("utf-8", "ignore").decode("utf-8"))

if __name__ == "__main__":
    sys.stdout.reconfigure(encoding='utf-8')
    extract_text_from_ppt(sys.argv[1])
